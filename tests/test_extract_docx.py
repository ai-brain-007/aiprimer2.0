"""Tests for pipeline.extract.docx (fixtures generated with python-docx at test time)."""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import docx
import pytest

from pipeline.config import Settings, load_settings
from pipeline.extract import extract_file
from pipeline.extract.docx import docx_to_markdown_python, extract_docx

REPO_ROOT = Path("/home/user/aiprimer2.0")


@pytest.fixture
def settings() -> Settings:
    return load_settings(repo_root=REPO_ROOT)


def make_docx(path: Path, with_props: bool = True) -> Path:
    d = docx.Document()
    d.add_heading("Test Document", level=1)
    d.add_paragraph("Hello world paragraph with café.")
    d.add_heading("Second Part", level=2)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "42"
    d.add_paragraph("Bullet item", style="List Bullet")
    d.add_paragraph("Closing paragraph.")
    if with_props:
        d.core_properties.author = "Jane Doe"
        d.core_properties.title = "Docx Title"
        d.core_properties.created = datetime(2020, 5, 17, 10, 0, 0)
    else:
        d.core_properties.author = ""
        d.core_properties.title = ""
    d.save(str(path))
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    return make_docx(tmp_path / "doc.docx")


def _assert_content(text: str) -> None:
    assert "# Test Document" in text
    assert "## Second Part" in text
    assert "Hello world paragraph with café." in text
    assert "| name" in text and "| alpha" in text and "42" in text
    assert "Bullet item" in text
    assert text.index("# Test Document") < text.index("Hello world") < text.index("| name") < text.index("Closing")
    assert text.endswith("\n") and "\r" not in text


def test_extract_docx_content_and_metadata(docx_file: Path, settings: Settings, tmp_path: Path):
    ex = extract_docx(docx_file, settings, tmp_path / "work")
    assert ex.source_type == "docx"
    assert ex.transcript_kind == "text"
    assert ex.language == ""
    assert ex.needs_vision is False
    _assert_content(ex.text)
    md = ex.metadata
    assert md.title == "Docx Title"
    assert md.author_raw == "Jane Doe"
    assert md.published_date == "2020-05-17"
    assert md.date_precision == "day"
    assert 0.3 <= md.confidence <= 0.8
    assert "title from docx core properties" in md.evidence
    assert "author from docx core properties" in md.evidence
    if shutil.which("pandoc"):
        assert not any("pandoc" in w for w in ex.warnings)


def test_python_docx_fallback_when_pandoc_missing(docx_file: Path, settings: Settings, tmp_path: Path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda *a, **k: None)
    ex = extract_docx(docx_file, settings, tmp_path / "work")
    assert any("pandoc not found" in w for w in ex.warnings)
    _assert_content(ex.text)
    assert ex.metadata.author_raw == "Jane Doe"


def test_python_docx_renderer_headings_and_tables(docx_file: Path):
    md = docx_to_markdown_python(docx_file)
    assert md.startswith("# Test Document")
    assert "| name | value |" in md
    assert "| --- | --- |" in md
    assert "| alpha | 42 |" in md
    assert "- Bullet item" in md


def test_title_falls_back_to_first_heading(settings: Settings, tmp_path: Path):
    path = make_docx(tmp_path / "notitle.docx", with_props=False)
    ex = extract_docx(path, settings, tmp_path / "work")
    assert ex.metadata.title == "Test Document"
    assert "title from first heading" in ex.metadata.evidence
    assert ex.metadata.author_raw == ""


def test_extract_file_dispatches_docx(docx_file: Path, settings: Settings, tmp_path: Path):
    ex = extract_file(docx_file, "docx", settings, tmp_path / "work")
    assert ex.source_type == "docx"
    assert "# Test Document" in ex.text
